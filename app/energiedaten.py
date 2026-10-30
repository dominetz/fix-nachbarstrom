"""Energiedaten aus dem EDA-Portal: Excel-Export lesen und schreiben.

Das EDA-Portal liefert die Energiedaten je Monat als Excel-Datei. Der genaue
Aufbau ist nicht öffentlich dokumentiert. Belegt sind die Spaltennamen
„Eigendeckung gemeinschaftlicher Erzeugung [kWh]“ (Abnehmer),
„Gesamte gemeinschaftliche Erzeugung [kWh]“ und „Restüberschuss bei EG und je
ZP [kWh]“ (Erzeuger), Quelle: kleinkraft.co.at, 05.03.2026. Der Rest ist
nachgebaut. Das Einlesen sucht deshalb Zeilen und Spalten über Stichworte statt
über feste Zellen und lässt sich an den ersten echten Export anpassen.

Aufbau des Blatts „Energiedaten“ (eine Spalte je Zählpunkt und Messgröße):

    MeteringpointID         | AT…01        | AT…01        | AT…02 …
    Energierichtung         | CONSUMPTION  | CONSUMPTION  | GENERATION
    Beginn der Messperiode  | Gesamtverbrauch lt. Messung [kWh] | Eigendeckung … [kWh] | …
    01.10.2025 00:00:00     | 0,123        | 0,000        | …
"""
import io
import math
import random
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook

from .models import Profiltyp, Teilnehmer
from .zaehlpunkt import normalisieren

VERBRAUCH = "Gesamtverbrauch lt. Messung [kWh]"
ANTEIL = "Anteil gemeinschaftliche Erzeugung [kWh]"
EIGENDECKUNG = "Eigendeckung gemeinschaftlicher Erzeugung [kWh]"
ERZEUGUNG = "Gesamte gemeinschaftliche Erzeugung [kWh]"
RESTUEBERSCHUSS = "Restüberschuss bei EG und je ZP [kWh]"

SPALTEN_ABNEHMER = [(VERBRAUCH, "verbrauch"), (ANTEIL, "anteil"), (EIGENDECKUNG, "eigendeckung")]
SPALTEN_ERZEUGER = [(ERZEUGUNG, "erzeugung"), (RESTUEBERSCHUSS, "restueberschuss")]
FELDER = ("verbrauch", "anteil", "eigendeckung", "erzeugung", "restueberschuss")

# Stichworte zum Erkennen der Spalten, die ersten Treffer gewinnen
_STICHWORTE = [("eigendeckung", "eigendeckung"), ("restüberschuss", "restueberschuss"),
               ("restueberschuss", "restueberschuss"), ("anteil", "anteil"),
               ("erzeugung", "erzeugung"), ("verbrauch", "verbrauch")]

ZEITFORMAT = "%Y-%m-%d %H:%M"


def viertelstunden(von: date, bis: date) -> list[datetime]:
    """Alle Viertelstunden von 00:00 des ersten bis 23:45 des letzten Tages."""
    start = datetime(von.year, von.month, von.day)
    anzahl = ((bis - von).days + 1) * 96
    return [start + timedelta(minutes=15 * i) for i in range(anzahl)]


# ---------------------------------------------------------------- Schreiben

@dataclass
class Spur:
    """Werte eines Zählpunkts, je Viertelstunde."""
    teilnehmer: Teilnehmer
    werte: dict[str, list[float]] = field(default_factory=dict)

    @property
    def erzeuger(self) -> bool:
        return self.teilnehmer.rolle.value == "erzeuger"


def export_schreiben(spuren: list[Spur], zeiten: list[datetime], gemeinschafts_id: str) -> bytes:
    """Excel-Datei im Aufbau des EDA-Exports (ein Monat)."""
    wb = Workbook()
    ue = wb.active
    ue.title = "Übersicht"
    von, bis = zeiten[0].date(), zeiten[-1].date()
    ue.append([f"Energiedaten Gemeinschaft {gemeinschafts_id}"])
    ue.append([f"Zeitraum: {von:%d.%m.%Y} – {bis:%d.%m.%Y}"])
    verfuegbar = sum(sum(s.werte.get("erzeugung", [])) for s in spuren)
    ue.append(["Gesamte der Gemeinschaft zur Verfügung stehende Energiemenge [kWh]", round(verfuegbar, 3)])
    ue.append([])
    ue.append(["Zählpunkt", "Energierichtung", "Name", VERBRAUCH, ANTEIL, EIGENDECKUNG, ERZEUGUNG, RESTUEBERSCHUSS])
    for s in spuren:
        summe = {f: round(sum(s.werte.get(f, [])), 3) for f in FELDER}
        ue.append([s.teilnehmer.zaehlpunkt, "GENERATION" if s.erzeuger else "CONSUMPTION", s.teilnehmer.name,
                   *(summe[f] if f in s.werte else None for f in FELDER)])

    ws = wb.create_sheet("Energiedaten")
    kopf_zp, kopf_richtung, kopf_name, spalten = ["MeteringpointID"], ["Energierichtung"], ["Beginn der Messperiode"], []
    for s in spuren:
        for titel, feld in (SPALTEN_ERZEUGER if s.erzeuger else SPALTEN_ABNEHMER):
            kopf_zp.append(s.teilnehmer.zaehlpunkt)
            kopf_richtung.append("GENERATION" if s.erzeuger else "CONSUMPTION")
            kopf_name.append(titel)
            spalten.append(s.werte[feld])
    ws.append(kopf_zp)
    ws.append(kopf_richtung)
    ws.append(kopf_name)
    for i, zeit in enumerate(zeiten):
        ws.append([zeit.strftime("%d.%m.%Y %H:%M:%S"), *(round(sp[i], 3) for sp in spalten)])
    ws.column_dimensions["A"].width = 22
    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


# ---------------------------------------------------------------- Einlesen

@dataclass
class Import:
    werte: list[dict]  # je Zeile: zaehlpunkt, beginn, und die Felder
    zaehlpunkte: set[str]
    von: datetime
    bis: datetime


def _zeit(wert) -> datetime | None:
    if isinstance(wert, datetime):
        return wert
    if not isinstance(wert, str):
        return None
    text = wert.strip()
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _zahl(wert) -> float:
    if wert is None or wert == "":
        return 0.0
    if isinstance(wert, (int, float)):
        return float(wert)
    text = str(wert).strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "")
    return float(text.replace(",", "."))


def _feld(titel: str) -> str | None:
    t = (titel or "").lower()
    for wort, feld in _STICHWORTE:
        if wort in t:
            return feld
    return None


def export_einlesen(inhalt: bytes) -> Import:
    try:
        wb = load_workbook(io.BytesIO(inhalt), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - beliebige kaputte Datei
        raise ValueError("Die Datei ist keine lesbare Excel-Datei (.xlsx).") from e
    for ws in wb.worksheets:
        zeilen = ws.iter_rows(values_only=True)
        kopf = []
        for zeile in zeilen:
            if zeile and _zeit(zeile[0]):
                return _werte_lesen(kopf, zeile, zeilen)
            kopf.append(zeile)
            if len(kopf) > 15:
                break
    raise ValueError("In der Datei wurde kein Blatt mit Viertelstundenwerten gefunden.")


def _werte_lesen(kopf: list[tuple], erste: tuple, rest) -> Import:
    zp_zeile = next((z for z in kopf if z and str(z[0] or "").lower().startswith(("meteringpoint", "zählpunkt",
                                                                                   "zaehlpunkt"))), None)
    titel_zeile = kopf[-1] if kopf else None
    if not zp_zeile or not titel_zeile:
        raise ValueError("Kopfzeilen mit Zählpunkt und Messgröße fehlen.")
    spalten = []  # (index, zählpunkt, feld)
    for i in range(1, len(zp_zeile)):
        zp, feld = normalisieren(str(zp_zeile[i] or "")), _feld(str(titel_zeile[i] if i < len(titel_zeile) else ""))
        if zp and feld:
            spalten.append((i, zp, feld))
    if not spalten:
        raise ValueError("Keine bekannten Spalten gefunden (z. B. „Eigendeckung“, „Erzeugung“).")
    zeilen: dict[tuple[str, str], dict] = {}
    von = bis = None
    for zeile in [erste, *rest]:
        zeit = _zeit(zeile[0]) if zeile else None
        if not zeit:
            continue
        von, bis = von or zeit, zeit
        beginn = zeit.strftime(ZEITFORMAT)
        for i, zp, feld in spalten:
            ziel = zeilen.setdefault((zp, beginn), {"zaehlpunkt": zp, "beginn": beginn, **dict.fromkeys(FELDER, 0.0)})
            ziel[feld] += _zahl(zeile[i] if i < len(zeile) else None)
    return Import(list(zeilen.values()), {zp for _, zp, _ in spalten}, von, bis)


# ---------------------------------------------------------------- Fiktive Daten

# Annahme Vorarlberg: Anteil des PV-Jahresertrags je Monat (Jan…Dez), Summe 100 %
PV_MONAT = [3.0, 5.0, 8.5, 11.0, 12.5, 13.0, 13.5, 12.0, 9.0, 6.5, 3.5, 2.5]
# Sonnenauf- und -untergang Mitte des Monats, Ortszeit, gerundet
SONNE = [(8.0, 17.0), (7.4, 17.8), (6.6, 18.5), (6.6, 20.2), (5.8, 20.9), (5.4, 21.3),
         (5.7, 21.2), (6.3, 20.5), (7.0, 19.5), (7.7, 18.5), (7.5, 16.9), (8.0, 16.6)]


def _jahreszeit(monat: int) -> str:
    return "winter" if monat in (11, 12, 1, 2) else "sommer" if monat in (5, 6, 7, 8) else "uebergang"


def _lastform(typ: Profiltyp, tag: date) -> list[float]:
    from .matching import _form_verbrauch
    form = _form_verbrauch(typ, _jahreszeit(tag.month))
    if typ == Profiltyp.GEWERBE:  # Betrieb nur Mo–Fr
        offen = tag.weekday() < 5
        form = [0.25 + (1.6 if offen and 7.5 <= (q + 0.5) / 4 <= 17.5 else 0.0) for q in range(96)]
    return form


def simulieren(erzeuger: Teilnehmer, abnehmer: list[Teilnehmer], zeiten: list[datetime],
               seed: int = 1) -> list[Spur]:
    """Fiktive Viertelstundenwerte mit Wetter. Geteilt wird dynamisch nach Verbrauch."""
    zufall = random.Random(seed)
    tage = sorted({z.date() for z in zeiten})
    wetter = {t: zufall.betavariate(2.2, 1.3) for t in tage}  # 0 trüb … 1 klar
    mittel_wetter = 2.2 / 3.5

    def last(t: Teilnehmer, jahres_kwh: float) -> list[float]:
        formen = {}
        werte = []
        for z in zeiten:
            if z.date() not in formen:
                formen[z.date()] = _lastform(t.profiltyp, z.date())
            werte.append(formen[z.date()][z.hour * 4 + z.minute // 15])
        # auf den Jahresverbrauch skalieren, anteilig für den Zeitraum
        faktor = jahres_kwh * len(zeiten) / 35040 / (sum(werte) or 1)
        return [max(0.0, w * faktor * zufall.uniform(0.8, 1.2)) for w in werte]

    pv = []
    jahr = erzeuger.anlage_kwp * 1050
    for z in zeiten:
        auf, unter = SONNE[z.month - 1]
        stunde = z.hour + z.minute / 60 + 0.125
        if not auf < stunde < unter:
            pv.append(0.0)
            continue
        tage_im_monat = 31 if z.month in (1, 3, 5, 7, 8, 10, 12) else 30 if z.month != 2 else 28
        tag_kwh = jahr * PV_MONAT[z.month - 1] / 100 / tage_im_monat
        # Glockenform sin², Summe der Viertelstunden über den Tag = 1
        form = math.sin(math.pi * (stunde - auf) / (unter - auf)) ** 2 * 0.5 / (unter - auf)
        pv.append(max(0.0, tag_kwh * form * wetter[z.date()] / mittel_wetter * zufall.uniform(0.9, 1.1)))

    eigen = last(erzeuger, erzeuger.jahresverbrauch_kwh)
    lasten = {a.id: last(a, a.jahresverbrauch_kwh) for a in abnehmer}
    einspeisung = [max(0.0, p - e) for p, e in zip(pv, eigen)]
    spuren = {a.id: Spur(a, {"verbrauch": lasten[a.id], "anteil": [], "eigendeckung": []}) for a in abnehmer}
    rest = []
    for i, ueberschuss in enumerate(einspeisung):
        bedarf = sum(lasten[a.id][i] for a in abnehmer)
        verteilt = 0.0
        for a in abnehmer:
            anteil = ueberschuss * lasten[a.id][i] / bedarf if bedarf else 0.0
            deckung = min(anteil, lasten[a.id][i])
            spuren[a.id].werte["anteil"].append(deckung)
            spuren[a.id].werte["eigendeckung"].append(deckung)
            verteilt += deckung
        rest.append(ueberschuss - verteilt)
    return [Spur(erzeuger, {"erzeugung": einspeisung, "restueberschuss": rest}), *spuren.values()]


def monate(von: date, bis: date) -> list[tuple[date, date]]:
    """Kalendermonate zwischen zwei Tagen, je (erster, letzter Tag)."""
    liste, tag = [], date(von.year, von.month, 1)
    while tag <= bis:
        naechster = date(tag.year + tag.month // 12, tag.month % 12 + 1, 1)
        liste.append((tag, naechster - timedelta(days=1)))
        tag = naechster
    return liste
